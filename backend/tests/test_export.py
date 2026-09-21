from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import load_workbook

from app.export import build_leads_workbook
from app.models import Lead, LeadStatus


def _make_lead(**overrides) -> Lead:
    now = datetime.now(timezone.utc)
    defaults = dict(
        id="lead-1",
        job_id="job-1",
        source_file="card1.jpg",
        first_name="Jane",
        last_name="Doe",
        job_title="CEO",
        company="Acme Corp",
        location="Austin, TX",
        phone="+15125551234",
        email="jane@acme.com",
        confidence=0.9,
        status=LeadStatus.DONE,
        created_at=now,
        updated_at=now,
    )
    defaults.update(overrides)
    return Lead(**defaults)


def test_export_produces_valid_workbook():
    leads = [_make_lead()]
    xlsx_bytes = build_leads_workbook(leads)
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert "Leads" in wb.sheetnames
    ws = wb["Leads"]
    header = [c.value for c in ws[1]]
    assert "First Name" in header
    assert "Phone Number" in header
    assert ws[2][0].value == "Jane"


def test_phone_column_is_text_formatted_to_avoid_scientific_notation():
    leads = [_make_lead(phone="+15125551234")]
    xlsx_bytes = build_leads_workbook(leads)
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Leads"]
    header = [c.value for c in ws[1]]
    phone_col_idx = header.index("Phone Number") + 1
    cell = ws.cell(row=2, column=phone_col_idx)
    assert cell.number_format == "@"
    assert cell.value == "+15125551234"


def test_excludes_non_done_leads():
    leads = [_make_lead(status=LeadStatus.DONE), _make_lead(id="lead-2", status=LeadStatus.FAILED)]
    xlsx_bytes = build_leads_workbook(leads)
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Leads"]
    assert ws.max_row == 2  # header + 1 done lead


def test_summary_sheet_counts_failed():
    leads = [_make_lead(status=LeadStatus.DONE), _make_lead(id="lead-2", status=LeadStatus.FAILED)]
    xlsx_bytes = build_leads_workbook(leads)
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    summary = wb["Summary"]
    rows = {row[0].value: row[1].value for row in summary.iter_rows()}
    assert rows["Cards that failed extraction"] == 1
    assert rows["Total cards uploaded"] == 2


def test_empty_leads_list_does_not_crash():
    xlsx_bytes = build_leads_workbook([])
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert wb["Leads"].max_row == 1  # header only
